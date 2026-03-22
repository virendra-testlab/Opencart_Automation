import openpyxl

def getTotalRows(File_path, sheet):
    workbook = openpyxl.load_workbook(File_path)
    Sheet = workbook[sheet]
    return Sheet.max_row

def getTotalColumns(file_path, sheet):
    workbook = openpyxl.load_workbook(file_path)
    Sheet = workbook[sheet]
    return Sheet.max_column

def getReadData(file, sheet, rows, col):
    workbook = openpyxl.load_workbook(file)
    R_Sheet = workbook[sheet]
    return R_Sheet.cell(rows, col).value

def writeData(file, sheet, rows, col, data):
    workbook = openpyxl.load_workbook(file)
    W_sheet = workbook[sheet]
    W_sheet.cell(rows, col).value = data
    workbook.save(file)